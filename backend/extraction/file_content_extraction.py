"""Extract plain text from a resume file (PDF, DOCX, or plain text).

This is the only module that should know how a resume — or a plain-text
job description, such as the fixtures under tests/fixtures/ — is
structured on disk. Everything downstream (matching, PII handling) works
with plain text and doesn't care what the original file format was.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class UnsupportedResumeFormat(ValueError):
    """Raised when the file isn't a format we know how to read."""


def extract_text(file_path: str | Path) -> str:
    """Extract plain text from a resume or job-description file.

    Dispatches on the file extension: `.pdf` goes through pypdf, `.docx`
    goes through python-docx, `.txt` is read directly (used for plain-text
    job description fixtures — no parsing needed).

    Raises:
        FileNotFoundError: `file_path` doesn't exist.
        UnsupportedResumeFormat: the extension isn't `.pdf`, `.docx`, or `.txt`.
        ValueError: the file was read but no text could be extracted
            (e.g. a scanned/image-only PDF with no text layer).
    """
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf_text(path)
    if suffix == ".docx":
        return _extract_docx_text(path)
    if suffix == ".txt":
        return _extract_txt_text(path)

    raise UnsupportedResumeFormat(
        f"Unsupported format {suffix!r}. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
    )


def _extract_pdf_text(path: Path) -> str:
    reader = PdfReader(path)
    pages_text = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages_text).strip()

    if not text:
        raise ValueError(
            f"No extractable text found in {path.name} — it may be a "
            "scanned/image-only PDF with no text layer."
        )
    return text


def _extract_docx_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # `.paragraphs` skips table cells, but resumes often lay out skills or
    # experience in tables — walk those separately so nothing is dropped.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    text = "\n".join(paragraphs).strip()
    if not text:
        raise ValueError(f"No extractable text found in {path.name}.")
    return text


def _extract_txt_text(path: Path) -> str:
    text = path.read_text(encoding="utf-8").strip()
    if not text:
        raise ValueError(f"No extractable text found in {path.name}.")
    return text
